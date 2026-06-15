"""Shared document-building blocks used by the TRP/VTR/BRP generators."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches

from apps.frameworks.models import ClauseStatus

from .docx_utils import (
    add_rtl_paragraph,
    build_header_footer,
    format_run,
    make_table,
    set_cell_text,
    set_page_letter,
    set_paragraph_rtl,
    set_section_rtl,
    shade_cell,
    shade_paragraph,
    shamsi_date,
)

# Evidence attachments embedded as images: only these content types are
# pictures python-docx can render inline.
IMAGE_CONTENT_TYPES = {"image/png", "image/jpeg"}

# [[filename_slug]] placeholders inside clause text, resolved against the
# clause's image library (apps.assessments.models.Attachment).
IMAGE_TOKEN_RE = re.compile(r"\[\[([^\[\]\r\n]+)\]\]")
# Fits inside the ~75%-width value column of the per-clause table
# (6.5in content width minus margins, times 0.75).
EVIDENCE_IMAGE_WIDTH = Inches(4.5)

LAB_NAME = "مرکز ارزیابی ایمنی و امنیتی تبادل امن"

LOGO_PATH = Path(__file__).parent / "assets" / "lab_logo.png"
LOGO_WIDTH = Inches(1.8)

STATUS_RESULT_TEXT = {
    ClauseStatus.COMPLIANT: "قبول",
    ClauseStatus.FINDING: "عدم انطباق",
    ClauseStatus.NOT_APPLICABLE: "مصداق ندارد",
    ClauseStatus.UNREVIEWED: "",
}


def result_text_and_highlight(status: str) -> tuple[str, bool]:
    text = STATUS_RESULT_TEXT.get(status, "")
    highlight = status == ClauseStatus.FINDING
    return text, highlight


def new_document(*, doc_title: str, assessment, doc_code_prefix: str) -> Document:
    document = Document()
    set_page_letter(document)
    set_section_rtl(document)
    system = assessment.system
    build_header_footer(
        document,
        doc_title=doc_title,
        product=system.name,
        company=system.company.name,
        doc_code=f"{doc_code_prefix}-{system.name}",
    )
    return document


def add_cover(document, *, doc_title: str, assessment, logo: bool = False) -> None:
    system = assessment.system
    add_rtl_paragraph(document, "به نام خدا", size=14, bold=True,
                      align=WD_ALIGN_PARAGRAPH.CENTER)
    if logo and LOGO_PATH.exists():
        document.add_picture(str(LOGO_PATH), width=LOGO_WIDTH)
        document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_rtl_paragraph(document, "", align=WD_ALIGN_PARAGRAPH.CENTER)
    add_rtl_paragraph(document, doc_title, size=20, bold=True,
                      align=WD_ALIGN_PARAGRAPH.CENTER)
    add_rtl_paragraph(document, "", align=WD_ALIGN_PARAGRAPH.CENTER)
    add_rtl_paragraph(document, f"نام محصول: {system.name}", size=13,
                      align=WD_ALIGN_PARAGRAPH.CENTER)
    add_rtl_paragraph(document, f"نام شرکت: {system.company.name}", size=13,
                      align=WD_ALIGN_PARAGRAPH.CENTER)
    add_rtl_paragraph(document, "نام آزمایشگاه:", size=13,
                      align=WD_ALIGN_PARAGRAPH.CENTER)
    add_rtl_paragraph(document, LAB_NAME, size=13, bold=True,
                      align=WD_ALIGN_PARAGRAPH.CENTER)
    add_rtl_paragraph(document, f"نسخه سند: {assessment.doc_version}", size=11,
                      align=WD_ALIGN_PARAGRAPH.CENTER)
    document.add_page_break()


def add_change_log_table(document, assessment, *, min_rows: int = 8) -> None:
    add_rtl_paragraph(document, "تغییرات سند", size=13, bold=True,
                      align=WD_ALIGN_PARAGRAPH.CENTER)
    entries = list(assessment.change_log or [])
    n_rows = max(min_rows, len(entries)) + 1
    table = make_table(document, rows=n_rows, cols=4)
    headers = ["ردیف", "نسخه", "تاریخ", "شرح تغییرات"]
    for col, text in enumerate(headers):
        set_cell_text(table.rows[0].cells[col], text, bold=True,
                      align=WD_ALIGN_PARAGRAPH.CENTER)
        shade_cell(table.rows[0].cells[col])
    for idx in range(1, n_rows):
        entry = entries[idx - 1] if idx - 1 < len(entries) else {}
        set_cell_text(table.rows[idx].cells[0], str(idx),
                      align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(table.rows[idx].cells[1], str(entry.get("version", "")),
                      align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(table.rows[idx].cells[2], str(entry.get("date", "")),
                      align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(table.rows[idx].cells[3], str(entry.get("description", "")))
    document.add_page_break()


def add_evaluation_specs(document, assessment, *, table_caption: str) -> None:
    """Section 1: مشخصات ارزیابی — tester/approver codes, completion date."""
    add_rtl_paragraph(document, table_caption, size=11)
    table = make_table(document, rows=3, cols=2)
    completed = (
        shamsi_date(assessment.test_completed_date)
        if assessment.test_completed_date
        else ""
    )
    rows = [
        ("کد آزمونگر", assessment.tester_code),
        ("کد تأییدکننده", assessment.approver_code),
        ("تاریخ اتمام آزمون", completed),
    ]
    for idx, (label, value) in enumerate(rows):
        set_cell_text(table.rows[idx].cells[0], label, bold=True)
        shade_cell(table.rows[idx].cells[0])
        set_cell_text(table.rows[idx].cells[1], value)


def add_callout(document, text: str, *, fill: str = "C6E0B4") -> None:
    """A shaded note paragraph (e.g. the green 'توضیحات' box)."""
    paragraph = add_rtl_paragraph(document, text, size=10)
    shade_paragraph(paragraph, fill)


def add_diagram_placeholder(document, *, height: Inches = Inches(3)) -> None:
    """An empty bordered area where the assessor manually inserts a diagram
    (network/architecture) directly in Word."""
    table = make_table(document, rows=1, cols=1)
    cell = table.rows[0].cells[0]
    table.rows[0].height = height
    set_cell_text(
        cell,
        "[محل قرارگیری نمودار — تصویر توسط ارزیاب درج می‌شود]",
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )


def add_clause_evidence_images(cell, clause_assessment) -> None:
    """Embed each image attachment from the clause's sub-clause evidence,
    centered below the existing cell content."""
    for sub_assessment in clause_assessment.sub_assessments.all():
        for attachment in sub_assessment.attachments.all():
            if attachment.content_type not in IMAGE_CONTENT_TYPES:
                continue
            path = Path(attachment.file.path)
            if not path.exists():
                continue
            paragraph = cell.add_paragraph()
            set_paragraph_rtl(paragraph)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run()
            run.add_picture(str(path), width=EVIDENCE_IMAGE_WIDTH)


def _insert_clause_image_or_warning(paragraph, token, clause_assessment, *, size) -> None:
    """Resolve a [[token]] placeholder to an inline image from the image
    owner's library for this clause, or render a visible warning. Bad
    tokens (missing image, deleted file, corrupt image) never raise."""
    try:
        from apps.assessments.models import Attachment

        owner_id = clause_assessment.updated_by_id or clause_assessment.assessment.assessor_id
        attachment = Attachment.objects.filter(
            clause_assessment=clause_assessment,
            uploaded_by_id=owner_id,
            filename_slug=token,
        ).first()
        if attachment is None or attachment.content_type not in IMAGE_CONTENT_TYPES:
            raise FileNotFoundError(token)
        path = Path(attachment.file.path)
        if not path.exists():
            raise FileNotFoundError(token)
        run = paragraph.add_run()
        run.add_picture(str(path), width=EVIDENCE_IMAGE_WIDTH)
    except Exception:
        run = paragraph.add_run(f"[تصویر یافت نشد: {token}]")
        format_run(run, size=size, bold=True, highlight=True)


def set_cell_text_with_images(cell, text, clause_assessment, *, size=11,
                               highlight=False, bold=False,
                               align=WD_ALIGN_PARAGRAPH.RIGHT) -> None:
    """Like set_cell_text, but expands [[filename_slug]] placeholders into
    inline images (or a visible warning) at their exact position in the
    text, with RTL preserved."""
    cell.text = ""
    first = True
    for line in (text or "").split("\n"):
        paragraph = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        set_paragraph_rtl(paragraph)
        paragraph.alignment = align

        pos = 0
        for match in IMAGE_TOKEN_RE.finditer(line):
            if match.start() > pos:
                run = paragraph.add_run(line[pos:match.start()])
                format_run(run, size=size, bold=bold, highlight=highlight)
            _insert_clause_image_or_warning(
                paragraph, match.group(1), clause_assessment, size=size
            )
            pos = match.end()
        if pos < len(line) or pos == 0:
            run = paragraph.add_run(line[pos:])
            format_run(run, size=size, bold=bold, highlight=highlight)


def add_clause_result_table(document, clause_assessment, *,
                            include_evidence: bool = False) -> None:
    """The per-clause 4-row table shared by TRP section 6 and the BRP."""
    clause = clause_assessment.clause
    result, highlight = result_text_and_highlight(clause_assessment.status)
    table = make_table(document, rows=4, cols=2)
    rows = [
        ("عنوان الزام", clause.description or f"{clause.code} {clause.title}", False),
        ("نتیجه نهایی آزمون", result, highlight),
        ("هدف الزام", clause.objective, False),
        ("تشریح آزمون انجام شده", clause_assessment.text, False),
    ]
    for idx, (label, value, value_highlight) in enumerate(rows):
        set_cell_text(table.rows[idx].cells[0], label, bold=True)
        shade_cell(table.rows[idx].cells[0])
        if idx == 3:
            set_cell_text_with_images(
                table.rows[idx].cells[1], value, clause_assessment,
                highlight=value_highlight, bold=value_highlight,
            )
        else:
            set_cell_text(table.rows[idx].cells[1], value, highlight=value_highlight,
                          bold=value_highlight)
    if include_evidence:
        add_clause_evidence_images(table.rows[3].cells[1], clause_assessment)
    # Label column ~25% width
    for row in table.rows:
        row.cells[0].width = document.sections[0].page_width * 1 // 4
        row.cells[1].width = document.sections[0].page_width * 3 // 4
