"""Generate the VTR (سند گزارش آزمون آسیب‌پذیری) Word document.

Section 2 is the per-category test-count table (rounds 1-4; only round 1
is populated in the MVP — multi-round support is backlog). Section 3 has
one table per OWASP category."""
from __future__ import annotations

import io
from itertools import groupby

from docx.enum.text import WD_ALIGN_PARAGRAPH

from apps.frameworks.models import ClauseStatus

from .common import (
    add_change_log_table,
    add_cover,
    add_evaluation_specs,
    new_document,
    result_text_and_highlight,
)
from .docx_utils import add_heading_fa, add_rtl_paragraph, make_table, \
    set_cell_text, shade_cell

DOC_TITLE = "سند گزارش آزمون آسیب‌پذیری"


def _grouped_by_category(assessment):
    cas = list(
        assessment.clause_assessments.select_related(
            "clause", "clause__requirement"
        ).order_by("clause__requirement__order", "clause__order")
    )
    return [
        (key, list(group))
        for key, group in groupby(cas, key=lambda ca: ca.clause.requirement)
    ]


def _category_result(items) -> str:
    """Aggregate verdict for one category."""
    statuses = {ca.status for ca in items}
    if ClauseStatus.FINDING in statuses:
        return ClauseStatus.FINDING
    if statuses == {ClauseStatus.NOT_APPLICABLE}:
        return ClauseStatus.NOT_APPLICABLE
    if ClauseStatus.UNREVIEWED in statuses:
        return ClauseStatus.UNREVIEWED
    return ClauseStatus.COMPLIANT


def generate_vtr(assessment) -> io.BytesIO:
    document = new_document(
        doc_title=DOC_TITLE, assessment=assessment, doc_code_prefix="VTR"
    )
    add_cover(document, doc_title=DOC_TITLE, assessment=assessment)
    add_change_log_table(document, assessment)

    grouped = _grouped_by_category(assessment)
    compliance = assessment.compliance_percent

    # 1 — مشخصات ارزیابی
    add_heading_fa(document, "1- مشخصات ارزیابی")
    add_evaluation_specs(document, assessment, table_caption="جدول 1-1 مشخصات آزمونگر")

    # 2 — گزارش ارزیابی آسیب‌پذیری
    add_heading_fa(document, "2- گزارش ارزیابی آسیب‌پذیری")
    add_rtl_paragraph(
        document,
        f"نتایج ارزیابی سامانه {assessment.system.name} در جدول زیر نشان داده "
        "شده است. جزئیات مربوط به هر آزمون در بخش تشریح آزمون‌های آسیب‌پذیری "
        "به تفصیل بیان شده است.",
        size=11,
    )
    add_rtl_paragraph(document, "جدول 2-1 گزارش تعداد آزمون‌های انجام شده", size=11)
    summary = make_table(document, rows=len(grouped) + 2, cols=5)
    headers = [
        "دسته آسیب‌پذیری",
        "آزمون مرتبه اول",
        "آزمون مرتبه دوم",
        "آزمون مرتبه سوم",
        "آزمون مرتبه چهارم",
    ]
    for col, text in enumerate(headers):
        cell = summary.rows[0].cells[col]
        set_cell_text(cell, text, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        shade_cell(cell)
    for idx, (requirement, items) in enumerate(grouped, start=1):
        cells = summary.rows[idx].cells
        tested = sum(
            1
            for ca in items
            if ca.status not in (ClauseStatus.UNREVIEWED, ClauseStatus.NOT_APPLICABLE)
        )
        set_cell_text(cells[0], requirement.klass_title)
        set_cell_text(cells[1], str(tested) if tested else "",
                      align=WD_ALIGN_PARAGRAPH.CENTER)
        for col in (2, 3, 4):  # rounds 2-4: multi-round support is backlog
            set_cell_text(cells[col], "", align=WD_ALIGN_PARAGRAPH.CENTER)
    final_row = summary.rows[len(grouped) + 1]
    final_row.cells[0].merge(final_row.cells[4])
    set_cell_text(
        final_row.cells[0],
        "نتیجه نهایی ارزیابی (درصد انطباق): "
        + (f"{compliance}٪" if compliance is not None else "-"),
        bold=True,
    )

    # 3 — تشریح آزمون‌های آسیب‌پذیری
    add_heading_fa(document, "3- تشریح آزمون‌های آسیب‌پذیری")
    for idx, (requirement, items) in enumerate(grouped, start=1):
        result, highlight = result_text_and_highlight(_category_result(items))
        table = make_table(document, rows=4, cols=2)
        otg_codes = "\n".join(ca.clause.code for ca in items)
        detail_parts = []
        for ca in items:
            detail_parts.append(f"{ca.clause.code}:")
            if ca.clause.objective:
                detail_parts.append(ca.clause.objective)
            detail_parts.append("نتیجه بررسی: " + (ca.text or ""))
            detail_parts.append("")
        rows = [
            ("عنوان آزمون", f"3-{idx} {requirement.klass_title}", False),
            ("هدف آزمون", "این آزمون مشتمل بر موارد زیر است:\n" + otg_codes, False),
            ("نتیجه آزمون", result, highlight),
            ("تشریح روال آزمون آسیب‌پذیری", "\n".join(detail_parts).strip(), False),
        ]
        for row_idx, (label, value, value_highlight) in enumerate(rows):
            set_cell_text(table.rows[row_idx].cells[0], label, bold=True)
            shade_cell(table.rows[row_idx].cells[0])
            set_cell_text(table.rows[row_idx].cells[1], value, highlight=value_highlight,
                          bold=value_highlight)
        add_rtl_paragraph(document, "")

    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer
