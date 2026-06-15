"""Low-level python-docx helpers for Persian RTL documents.

python-docx has no API for bidi/RTL/complex-script properties, so these
helpers manipulate the underlying WordprocessingML (oxml) directly. The
attribute values mirror the uploaded TRP/VTR sample documents:

- Persian (complex script) font: B Nazanin  — rFonts/@w:cs
- Latin font: Times New Roman               — rFonts/@w:ascii, @w:hAnsi
- RTL paragraphs: pPr/w:bidi
- RTL table column order: tblPr/w:bidiVisual
- Label cell shading: tcPr/w:shd/@w:fill = F2F2F2
- Finding result text: black text, rPr/w:highlight/@w:val = red
"""
from __future__ import annotations

import jdatetime
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

FA_FONT = "B Nazanin"
EN_FONT = "Times New Roman"
LABEL_SHADE = "F2F2F2"

# Palette mirroring the lab template
GROUP_LABEL_FILL = "1F3864"   # dark blue — info-table group labels, results header
FIELD_LABEL_FILL = "595959"   # dark gray — info-table field labels
HEADER_BLUE = "4472C4"        # table header rows
ROW_BLUE_LIGHT = "DDEBF7"
ROW_BLUE_DARK = "BDD7EE"
WHITE = "FFFFFF"
BLACK = "000000"

_FA_DIGITS = str.maketrans("0123456789", "۰۱۲۳۴۵۶۷۸۹")


def fa_num(value) -> str:
    """Render Latin digits as Persian numerals (۰۱۲…)."""
    return str(value).translate(_FA_DIGITS)


def _get_or_add(parent, tag: str):
    el = parent.find(qn(tag))
    if el is None:
        el = parent.makeelement(qn(tag), {})
        parent.append(el)
    return el


def set_paragraph_rtl(paragraph) -> None:
    """Mark a paragraph as right-to-left (w:bidi)."""
    pPr = paragraph._p.get_or_add_pPr()
    bidi = _get_or_add(pPr, "w:bidi")
    bidi.set(qn("w:val"), "1")


def format_run(
    run,
    *,
    size: int = 12,
    bold: bool = False,
    highlight: bool = False,
    color: str | None = None,
    fa_font: str = FA_FONT,
    en_font: str = EN_FONT,
) -> None:
    """Apply dual-script fonts, complex-script size and RTL to a run."""
    run.font.name = en_font
    run.font.size = Pt(size)
    run.font.bold = bold
    if highlight:
        run.font.highlight_color = WD_COLOR_INDEX.RED
    if color:
        run.font.color.rgb = RGBColor.from_string(color)

    rPr = run._r.get_or_add_rPr()
    rFonts = _get_or_add(rPr, "w:rFonts")
    rFonts.set(qn("w:ascii"), en_font)
    rFonts.set(qn("w:hAnsi"), en_font)
    rFonts.set(qn("w:cs"), fa_font)
    # complex-script size and bold mirror the latin values
    szCs = _get_or_add(rPr, "w:szCs")
    szCs.set(qn("w:val"), str(size * 2))
    if bold:
        _get_or_add(rPr, "w:bCs")
    rtl = _get_or_add(rPr, "w:rtl")
    rtl.set(qn("w:val"), "1")


def add_rtl_paragraph(
    container,
    text: str = "",
    *,
    size: int = 12,
    bold: bool = False,
    highlight: bool = False,
    color: str | None = None,
    align=WD_ALIGN_PARAGRAPH.RIGHT,
) -> object:
    """Add a fully RTL paragraph (works on document, cell, header, footer)."""
    paragraph = container.add_paragraph()
    set_paragraph_rtl(paragraph)
    paragraph.alignment = align
    if text:
        run = paragraph.add_run(text)
        format_run(run, size=size, bold=bold, highlight=highlight, color=color)
    return paragraph


def add_heading_fa(document, text: str, *, level: int = 1, size: int = 14,
                   page_break_before: bool = False):
    """Persian heading using the built-in Heading style (keeps navigation
    pane/TOC behaviour) with RTL + complex-script font overrides. Forced
    black (the lab template never uses blue headings)."""
    paragraph = document.add_heading(level=level)
    set_paragraph_rtl(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if page_break_before:
        paragraph.paragraph_format.page_break_before = True
    run = paragraph.add_run(text)
    format_run(run, size=size, bold=True, color=BLACK)
    return paragraph


def set_table_rtl(table) -> None:
    """Render table columns right-to-left (w:bidiVisual)."""
    tblPr = table._tbl.tblPr
    bidi = _get_or_add(tblPr, "w:bidiVisual")
    bidi.set(qn("w:val"), "1")
    table.alignment = WD_TABLE_ALIGNMENT.CENTER


def shade_cell(cell, fill: str = LABEL_SHADE) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = _get_or_add(tcPr, "w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)


def shade_paragraph(paragraph, fill: str) -> None:
    """Shade a paragraph's background (w:pPr/w:shd) — used for callout boxes."""
    pPr = paragraph._p.get_or_add_pPr()
    shd = _get_or_add(pPr, "w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)


def set_cell_text(
    cell,
    text: str,
    *,
    size: int = 11,
    bold: bool = False,
    highlight: bool = False,
    color: str | None = None,
    align=WD_ALIGN_PARAGRAPH.RIGHT,
) -> None:
    """Replace cell content with RTL text; '\n' becomes separate paragraphs."""
    cell.text = ""
    first = True
    for line in (text or "").split("\n"):
        paragraph = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        set_paragraph_rtl(paragraph)
        paragraph.alignment = align
        run = paragraph.add_run(line)
        format_run(run, size=size, bold=bold, highlight=highlight, color=color)


def make_table(document, rows: int, cols: int):
    table = document.add_table(rows=rows, cols=cols)
    table.style = "Table Grid"
    set_table_rtl(table)
    return table


def set_page_letter(document) -> None:
    """Letter page with 1-inch margins, matching the samples."""
    for section in document.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)


def set_section_rtl(document) -> None:
    """Mark each section as right-to-left (sectPr/w:bidi)."""
    for section in document.sections:
        bidi = _get_or_add(section._sectPr, "w:bidi")
        bidi.set(qn("w:val"), "1")


def shamsi_date(gregorian_date=None) -> str:
    """Format a date (or today) as a Shamsi yyyy/mm/dd string."""
    if gregorian_date is None:
        jdate = jdatetime.date.today()
    else:
        jdate = jdatetime.date.fromgregorian(date=gregorian_date)
    return jdate.strftime("%Y/%m/%d")


def build_header_footer(document, *, doc_title: str, product: str, company: str,
                        doc_code: str) -> None:
    """Page header: document title + product/company. Footer: date + code."""
    section = document.sections[0]
    header_p = section.header.paragraphs[0]
    set_paragraph_rtl(header_p)
    header_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = header_p.add_run(f"{doc_title} — {product} — {company}")
    format_run(run, size=9)

    footer_p = section.footer.paragraphs[0]
    set_paragraph_rtl(footer_p)
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer_p.add_run(f"تاریخ: {shamsi_date()}    |    {doc_code}")
    format_run(run, size=9)


def set_update_fields(document) -> None:
    """Tell Word to recompute fields (e.g. the TOC page numbers) on open."""
    settings = document.settings.element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)


def add_toc(document, *, levels: str = "1-2") -> None:
    """Insert a TOC field; Word fills in headings/page numbers on open
    (or after a manual right-click → Update Field)."""
    paragraph = document.add_paragraph()
    set_paragraph_rtl(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    begin_run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f' TOC \\o "{levels}" \\h \\z \\u '
    begin_run._r.append(fld_begin)
    begin_run._r.append(instr)

    separate_run = paragraph.add_run()
    fld_separate = OxmlElement("w:fldChar")
    fld_separate.set(qn("w:fldCharType"), "separate")
    separate_run._r.append(fld_separate)

    placeholder_run = paragraph.add_run(
        "برای به‌روزرسانی فهرست، کلیک راست کرده و گزینه Update Field را انتخاب کنید."
    )
    format_run(placeholder_run, size=10)

    end_run = paragraph.add_run()
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    end_run._r.append(fld_end)
