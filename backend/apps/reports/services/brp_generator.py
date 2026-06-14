"""Generate the BRP (بسته گزارش اشکالات): only the Finding clauses of a
TRP or VTR assessment, in the same per-clause table format."""
from __future__ import annotations

import io

from docx.enum.text import WD_ALIGN_PARAGRAPH

from apps.frameworks.models import ClauseStatus

from .common import add_clause_result_table, add_cover, new_document
from .docx_utils import add_heading_fa, add_rtl_paragraph

DOC_TITLE = "بسته گزارش اشکالات (BRP)"


def generate_brp(assessment) -> io.BytesIO:
    document = new_document(
        doc_title=DOC_TITLE, assessment=assessment, doc_code_prefix="BRP"
    )
    add_cover(document, doc_title=DOC_TITLE, assessment=assessment)

    findings = list(
        assessment.clause_assessments.select_related("clause", "clause__requirement")
        .prefetch_related("sub_assessments__attachments")
        .filter(status=ClauseStatus.FINDING)
        .order_by("clause__requirement__order", "clause__order")
    )

    add_heading_fa(document, "موارد عدم انطباق")
    add_rtl_paragraph(
        document,
        f"این سند شامل {len(findings)} مورد عدم انطباق شناسایی‌شده در ارزیابی "
        f"{assessment.get_kind_display()} سامانه {assessment.system.name} است.",
        size=11,
    )
    if not findings:
        add_rtl_paragraph(
            document,
            "در این ارزیابی هیچ مورد عدم انطباقی ثبت نشده است.",
            size=12,
            bold=True,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
    for idx, ca in enumerate(findings, start=1):
        add_heading_fa(
            document, f"{idx}- {ca.clause.title} ({ca.clause.code})", level=2, size=12
        )
        add_clause_result_table(document, ca, include_evidence=True)

    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer
