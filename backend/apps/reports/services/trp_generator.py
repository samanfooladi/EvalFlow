"""Generate the TRP (سند گزارش آزمون کارکردی) Word document, matching the
structure and RTL formatting of the lab's official template."""
from __future__ import annotations

import io
from itertools import groupby

from docx.enum.text import WD_ALIGN_PARAGRAPH

from apps.frameworks.models import ClauseStatus

from .common import (
    add_callout,
    add_change_log_table,
    add_clause_result_table,
    add_cover,
    add_diagram_placeholder,
    add_evaluation_specs,
    new_document,
    result_text_and_highlight,
)
from .docx_utils import add_heading_fa, add_rtl_paragraph, add_toc, make_table, \
    set_cell_text, set_update_fields, shade_cell

DOC_TITLE = "سند گزارش آزمون کارکردی"

RESULTS_NOTE = (
    "توضیحات: تغییرات زیرساختی، پیکربندی و نسخه سامانه مورد ارزیابی از ابتدا "
    "تا انتهای روند ارزیابی، در صورت وجود، باید توسط ارزیاب در این بخش ثبت شود."
)


def _ordered_clause_assessments(assessment):
    return list(
        assessment.clause_assessments.select_related(
            "clause", "clause__requirement"
        ).order_by("clause__requirement__order", "clause__order")
    )


def _class_summary_rows(clause_assessments):
    """Group by SFR class title preserving requirement order."""
    def key(ca):
        return ca.clause.requirement.klass_title

    rows = []
    for klass_title, group in groupby(clause_assessments, key=key):
        items = list(group)
        total = len(items)
        not_na = [ca for ca in items if ca.status != ClauseStatus.NOT_APPLICABLE]
        passed = sum(1 for ca in items if ca.status == ClauseStatus.COMPLIANT)
        percent = round(100 * passed / len(not_na)) if not_na else None
        rows.append(
            {
                "klass": klass_title,
                "total": total,
                "tested": len(not_na),
                "passed": passed,
                "percent": percent,
            }
        )
    return rows


def generate_trp(assessment) -> io.BytesIO:
    document = new_document(
        doc_title=DOC_TITLE, assessment=assessment, doc_code_prefix="TRP"
    )
    set_update_fields(document)
    add_cover(document, doc_title=DOC_TITLE, assessment=assessment, logo=True)

    add_rtl_paragraph(document, "فهرست", size=16, bold=True,
                      align=WD_ALIGN_PARAGRAPH.CENTER)
    add_toc(document)
    document.add_page_break()

    add_change_log_table(document, assessment)

    cas = _ordered_clause_assessments(assessment)
    system = assessment.system
    compliance = assessment.compliance_percent

    # 1 — مشخصات ارزیابی
    add_heading_fa(document, "1- مشخصات ارزیابی")
    add_evaluation_specs(document, assessment, table_caption="جدول 1-1 مشخصات آزمونگر")

    # 2 — معرفی محصول مورد ارزیابی
    add_heading_fa(document, "2- معرفی محصول مورد ارزیابی")
    add_rtl_paragraph(document, "جدول 2-1 شناسنامه محصول", size=11)
    profile = make_table(document, rows=4, cols=2)
    overall_result = ""
    if compliance is not None:
        overall_result = "قبول" if compliance == 100 else "عدم انطباق"
    profile_rows = [
        (
            "مشخصات کلی محصول",
            f"نام سامانه: {system.name}\n"
            f"نسخه محصول: {system.version}\n"
            f"شرکت تولیدکننده محصول: {system.company.name}\n"
            f"نام و شماره استانداردهای مورد ارزیابی: {assessment.framework.title}",
        ),
        ("مشخصات فنی محصول", system.description or ""),
        (
            "مشخصات ارزیابی",
            "نتیجه ارزیابی: " + overall_result + "\n"
            + (
                f"درصد انطباق: {compliance}٪"
                if compliance is not None
                else "درصد انطباق: -"
            ),
        ),
        ("قابلیت‌های کلان محصول", ""),
    ]
    for idx, (label, value) in enumerate(profile_rows):
        set_cell_text(profile.rows[idx].cells[0], label, bold=True)
        shade_cell(profile.rows[idx].cells[0])
        set_cell_text(profile.rows[idx].cells[1], value)

    # 3 — نمای کلی معماری محصول
    add_heading_fa(document, "3- نمای کلی معماری محصول")
    add_rtl_paragraph(document, assessment.architecture_overview or "", size=11)

    # 4 — پیکربندی آزمون
    add_heading_fa(document, "4- پیکربندی آزمون")
    add_rtl_paragraph(document, assessment.test_configuration or "", size=11)
    add_diagram_placeholder(document)

    # 5 — نتایج آزمون بر اساس کلاس‌های استاندارد معیار مشترک
    add_heading_fa(document, "5- نتایج آزمون بر اساس کلاس‌های استاندارد معیار مشترک")
    add_rtl_paragraph(document, "جدول 5-1 مشخصات آزمون‌های انجام شده حوزه SFRs", size=11)
    summary_rows = _class_summary_rows(cas)
    summary = make_table(document, rows=len(summary_rows) + 2, cols=6)
    headers = [
        "ردیف",
        "عنوان کلاس",
        "تعداد کل شاخص‌ها",
        "تعداد موارد آزمون شده",
        "تعداد آزمون‌های موفق",
        "درصد انطباق",
    ]
    for col, text in enumerate(headers):
        cell = summary.rows[0].cells[col]
        set_cell_text(cell, text, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        shade_cell(cell)
    for idx, row in enumerate(summary_rows, start=1):
        cells = summary.rows[idx].cells
        set_cell_text(cells[0], str(idx), align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(cells[1], row["klass"])
        set_cell_text(cells[2], str(row["total"]), align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(cells[3], str(row["tested"]), align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(cells[4], str(row["passed"]), align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(
            cells[5],
            f"{row['percent']}٪" if row["percent"] is not None else "-",
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
    final_row = summary.rows[len(summary_rows) + 1]
    final_row.cells[0].merge(final_row.cells[5])
    set_cell_text(
        final_row.cells[0],
        "نتیجه نهایی ارزیابی (درصد انطباق): "
        + (f"{compliance}٪" if compliance is not None else "-"),
        bold=True,
    )

    # 6 — نتایج و تشریح آزمون
    add_heading_fa(document, "6- نتایج و تشریح آزمون")
    results = make_table(document, rows=len(cas) + 1, cols=3)
    for col, text in enumerate(["ردیف", "عنوان الزام", "نتیجه آزمون"]):
        cell = results.rows[0].cells[col]
        set_cell_text(cell, text, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        shade_cell(cell)
    for idx, ca in enumerate(cas, start=1):
        cells = results.rows[idx].cells
        result, highlight = result_text_and_highlight(ca.status)
        set_cell_text(cells[0], str(idx), align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(cells[1], ca.clause.description or f"{ca.clause.code} {ca.clause.title}")
        set_cell_text(cells[2], result, highlight=highlight, bold=highlight,
                      align=WD_ALIGN_PARAGRAPH.CENTER)
    add_callout(document, RESULTS_NOTE)

    # 6-N per-clause subsections
    for idx, ca in enumerate(cas, start=1):
        add_heading_fa(
            document,
            f"6-{idx} {ca.clause.title} ({ca.clause.code})",
            level=2,
            size=12,
        )
        add_clause_result_table(document, ca)

    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer
