"""DOCX generation tests: round-trip the generated files through python-docx
and assert the RTL/formatting XML the lab's template requires is present."""
import base64

import pytest
from django.core.files.base import ContentFile
from django.core.management import call_command
from docx import Document

from apps.assessments.models import Assessment, Attachment
from apps.catalog.models import Company, ProductSystem
from apps.frameworks.models import ClauseStatus, Framework
from apps.reports.services.brp_generator import generate_brp
from apps.reports.services.trp_generator import generate_trp
from apps.reports.services.vtr_generator import generate_vtr

from .factories import AssessorFactory

# 1x1 transparent PNG
PNG_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY"
    "42YAAAAASUVORK5CYII="
)

pytestmark = pytest.mark.django_db

NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


@pytest.fixture(scope="module")
def frameworks(django_db_setup, django_db_blocker):
    with django_db_blocker.unblock():
        call_command("load_frameworks")


def _make_assessment(kind: str):
    framework = Framework.objects.get(kind=kind)
    company = Company.objects.create(name="شرکت آزمایشی پردازش")
    system = ProductSystem.objects.create(
        company=company, name="سامانه نمونه", version="1.0"
    )
    assessment = Assessment.objects.create(
        system=system,
        framework=framework,
        assessor=AssessorFactory(),
        tester_code="AS-01",
        approver_code="RV-02",
    )
    assessment.create_clause_assessments()
    # Verdicts: first clause a finding, one N/A, rest compliant.
    cas = list(assessment.clause_assessments.all())
    cas[0].apply_status(ClauseStatus.FINDING)
    cas[1].apply_status(ClauseStatus.NOT_APPLICABLE)
    for ca in cas[2:]:
        ca.apply_status(ClauseStatus.COMPLIANT)
    return assessment


@pytest.fixture
def trp_assessment(frameworks):
    return _make_assessment("TRP")


@pytest.fixture
def vtr_assessment(frameworks):
    return _make_assessment("VTR")


def _document_xml(buffer) -> str:
    from docx import Document as Open

    doc = Open(buffer)
    return doc.element.xml


def test_trp_generates_valid_docx(trp_assessment):
    buffer = generate_trp(trp_assessment)
    doc = Document(buffer)
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "به نام خدا" in text
    assert "سند گزارش آزمون کارکردی" in text
    assert "سامانه نمونه" in text
    assert "شرکت آزمایشی پردازش" in text


def test_trp_rtl_formatting_present(trp_assessment):
    xml = _document_xml(generate_trp(trp_assessment))
    assert "<w:bidi" in xml  # RTL paragraphs
    assert "<w:bidiVisual" in xml  # RTL tables
    assert "<w:rtl" in xml  # RTL runs
    assert 'w:cs="B Nazanin"' in xml  # Persian complex-script font
    assert 'w:fill="F2F2F2"' in xml  # label cell shading


def test_trp_finding_rendered_red(trp_assessment):
    xml = _document_xml(generate_trp(trp_assessment))
    assert 'w:val="FF0000"' in xml
    buffer = generate_trp(trp_assessment)
    doc = Document(buffer)
    all_text = "\n".join(
        cell.text for table in doc.tables for row in table.rows for cell in row.cells
    )
    assert "عدم انطباق" in all_text
    assert "قبول" in all_text
    assert "مصداق ندارد" in all_text


def test_trp_contains_class_summary_and_clause_tables(trp_assessment):
    doc = Document(generate_trp(trp_assessment))
    n_clauses = trp_assessment.clause_assessments.count()
    # change log + specs + profile + class summary + results list + per-clause
    assert len(doc.tables) >= n_clauses + 5
    all_text = "\n".join(p.text for p in doc.paragraphs)
    assert "نتایج آزمون بر اساس کلاس‌های استاندارد معیار مشترک" in all_text
    assert "کلاس ممیزی امنیت" in "\n".join(
        c.text for t in doc.tables for r in t.rows for c in r.cells
    )


def test_vtr_generates_with_categories(vtr_assessment):
    doc = Document(generate_vtr(vtr_assessment))
    cells = "\n".join(
        c.text for t in doc.tables for r in t.rows for c in r.cells
    )
    assert "آزمون جمع آوری اطلاعات" in cells
    assert "OTG-INFO-001" in cells
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "سند گزارش آزمون آسیب‌پذیری" in text


def test_vtr_rtl_formatting_present(vtr_assessment):
    xml = _document_xml(generate_vtr(vtr_assessment))
    for marker in ("<w:bidi", "<w:bidiVisual", 'w:cs="B Nazanin"'):
        assert marker in xml


def test_brp_contains_only_findings(trp_assessment):
    doc = Document(generate_brp(trp_assessment))
    cells = "\n".join(
        c.text for t in doc.tables for r in t.rows for c in r.cells
    )
    finding = trp_assessment.clause_assessments.filter(
        status=ClauseStatus.FINDING
    ).first()
    compliant = trp_assessment.clause_assessments.filter(
        status=ClauseStatus.COMPLIANT
    ).first()
    assert finding.clause.code in cells
    assert compliant.clause.code not in cells
    assert "قبول" not in cells  # no compliant verdicts anywhere


def test_brp_includes_evidence_images(trp_assessment):
    finding = trp_assessment.clause_assessments.filter(
        status=ClauseStatus.FINDING
    ).first()
    sub_assessment = finding.sub_assessments.first()
    Attachment.objects.create(
        assessment=trp_assessment,
        sub_clause_assessment=sub_assessment,
        file=ContentFile(PNG_1X1, name="evidence.png"),
        original_name="evidence.png",
        content_type="image/png",
        size=len(PNG_1X1),
    )
    doc = Document(generate_brp(trp_assessment))
    assert len(doc.inline_shapes) == 1


def test_brp_empty_when_no_findings(trp_assessment):
    trp_assessment.clause_assessments.filter(status=ClauseStatus.FINDING).update(
        status=ClauseStatus.COMPLIANT
    )
    doc = Document(generate_brp(trp_assessment))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "هیچ مورد عدم انطباقی ثبت نشده است" in text


def test_export_endpoint_streams_docx(api, as_user, trp_assessment):
    as_user(trp_assessment.assessor)
    response = api.get(f"/api/v1/assessments/{trp_assessment.pk}/export/?type=trp")
    assert response.status_code == 200
    assert response["Content-Type"].endswith("wordprocessingml.document")
    assert "attachment" in response["Content-Disposition"]


def test_export_blocked_when_unreviewed(api, as_user, trp_assessment):
    ca = trp_assessment.clause_assessments.first()
    ca.status = ClauseStatus.UNREVIEWED
    ca.save()
    as_user(trp_assessment.assessor)
    response = api.get(f"/api/v1/assessments/{trp_assessment.pk}/export/?type=trp")
    assert response.status_code == 400
    # BRP is always allowed (it only contains findings)
    response = api.get(f"/api/v1/assessments/{trp_assessment.pk}/export/?type=brp")
    assert response.status_code == 200


def test_export_allow_incomplete_only_for_elevated(api, as_user, trp_assessment, qa_lead):
    ca = trp_assessment.clause_assessments.first()
    ca.status = ClauseStatus.UNREVIEWED
    ca.save()
    as_user(trp_assessment.assessor)
    response = api.get(
        f"/api/v1/assessments/{trp_assessment.pk}/export/?type=trp&allow_incomplete=1"
    )
    assert response.status_code == 400  # assessor cannot override
    as_user(qa_lead)
    response = api.get(
        f"/api/v1/assessments/{trp_assessment.pk}/export/?type=trp&allow_incomplete=1"
    )
    assert response.status_code == 200


def test_export_wrong_type_for_kind(api, as_user, trp_assessment):
    as_user(trp_assessment.assessor)
    response = api.get(f"/api/v1/assessments/{trp_assessment.pk}/export/?type=vtr")
    assert response.status_code == 400


def test_export_denied_for_unassigned_assessor(api, as_user, trp_assessment):
    other = AssessorFactory()
    as_user(other)
    response = api.get(f"/api/v1/assessments/{trp_assessment.pk}/export/?type=trp")
    assert response.status_code == 403


def test_export_audited(api, as_user, trp_assessment):
    from apps.audit.models import AuditLog

    as_user(trp_assessment.assessor)
    api.get(f"/api/v1/assessments/{trp_assessment.pk}/export/?type=trp")
    assert AuditLog.objects.filter(action="export_docx").exists()
